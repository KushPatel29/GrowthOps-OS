import base64
import csv
import json
import re
import zlib
from pathlib import Path

import openpyxl

from growthops.export_bi import MARTS, PROJECT
from growthops.export_excel import build

ROOT = Path(__file__).resolve().parent.parent
CSV_DIR = ROOT / "dashboards" / "powerbi-data"


def test_every_power_bi_partition_embeds_the_committed_mart_exactly():
    model = (ROOT / PROJECT / "model.tmdl").read_text(encoding="utf-8")
    for mart in MARTS:
        assert f"ref table {mart}" in model, mart
        tmdl = (ROOT / PROJECT / "tables" / f"{mart}.tmdl").read_text(encoding="utf-8")
        encoded = re.search(r'Binary\.FromText\("([^"]+)"', tmdl).group(1)
        embedded = json.loads(zlib.decompress(base64.b64decode(encoded), -15))
        with (CSV_DIR / f"{mart}.csv").open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        assert embedded == rows[1:], mart
        for column in rows[0]:
            assert f"\tcolumn {column}\n" in tmdl, f"{mart}.{column}"


def test_excel_dashboard_is_formula_driven_and_audited(tmp_path):
    workbook = openpyxl.load_workbook(build(tmp_path / "dashboard.xlsx"))
    dashboard = workbook["Dashboard"]
    kpis = [dashboard.cell(row, 2).value for row in range(4, 15)]
    assert all(isinstance(value, str) and value.startswith("=") for value in kpis)
    audit = workbook["Audit"]
    checks = [audit.cell(row, 2).value for row in range(3, audit.max_row + 1) if audit.cell(row, 1).value]
    assert len(checks) >= 6 and all(str(value).startswith("=") for value in checks)
    assert all(value != "=0" for value in checks), "Audit checks must test source data"
    assert "SUMIFS(" in workbook["Period analysis"]["B9"].value
    committed = ROOT / "dashboards" / "GrowthOps_OS_Excel_Dashboard.xlsx"
    delivered = openpyxl.load_workbook(committed)
    assert delivered.sheetnames == workbook.sheetnames
    assert set(delivered.defined_names) == set(workbook.defined_names)
    for generated_sheet, delivered_sheet in zip(workbook.worksheets, delivered.worksheets):
        assert (delivered_sheet.max_row, delivered_sheet.max_column) == (
            generated_sheet.max_row, generated_sheet.max_column
        ), generated_sheet.title
        assert delivered_sheet.protection.sheet == generated_sheet.protection.sheet
        assert delivered_sheet.auto_filter.ref == generated_sheet.auto_filter.ref
        assert len(delivered_sheet._charts) == len(generated_sheet._charts)
        assert len(delivered_sheet.data_validations.dataValidation) == len(
            generated_sheet.data_validations.dataValidation
        )
        for generated_row, delivered_row in zip(generated_sheet, delivered_sheet):
            assert [(cell.value, cell.number_format) for cell in delivered_row] == [
                (cell.value, cell.number_format) for cell in generated_row
            ], generated_sheet.title
    cached = openpyxl.load_workbook(committed, data_only=True)
    assert cached["Dashboard"]["B4"].value == cached["Revenue"]["D2"].value / 100
    assert cached["Period analysis"]["B7"].value == "Ready"
    assert cached["Audit"]["B13"].value == "Yes"
    assert [cached["Audit"].cell(row, 2).value for row in range(3, 12)] == [0] * 9


def test_power_bi_model_has_date_dimension_relationships_and_documented_measures():
    from growthops.export_bi import DATE_KEYS, MEASURES

    definition = ROOT / PROJECT
    assert "ref table dim_date" in (definition / "model.tmdl").read_text(encoding="utf-8")
    date_table = (definition / "tables" / "dim_date.tmdl").read_text(encoding="utf-8")
    assert "partition dim_date = calculated" in date_table and "\t\tisKey\n" in date_table
    relationships = (definition / "relationships.tmdl").read_text(encoding="utf-8")
    for table, column in DATE_KEYS.items():
        assert f"fromColumn: {table}.{column}\n\ttoColumn: dim_date.Date" in relationships
    for table, measures in MEASURES.items():
        tmdl = (definition / "tables" / f"{table}.tmdl").read_text(encoding="utf-8")
        for name, _, fmt, folder, description in measures:
            block = tmdl.split(f"measure '{name}' = ", 1)
            assert len(block) == 2, f"{table}: {name}"
            assert f"/// {description}\n\t" in block[0][-len(description) - 12:], name
            assert f"formatString: {fmt}" in block[1][:400] and f"displayFolder: {folder}" in block[1][:400]


def test_excel_workbook_has_controls_and_marketing_kpis(tmp_path):
    workbook = openpyxl.load_workbook(build(tmp_path / "dashboard.xlsx"))
    assert {"Marketing KPIs", "Paid daily", "Email", "Links", "Definitions"} <= set(workbook.sheetnames)
    assert set(workbook.defined_names) >= {"StartDate", "EndDate", "FirstDataDate", "LastDataDate"}
    analysis = workbook["Period analysis"]
    assert {str(item.sqref) for item in analysis.data_validations.dataValidation} == {"B4", "B5"}
    assert "B5>=B4" in analysis.data_validations.dataValidation[1].formula1
    assert not analysis["B4"].protection.locked and analysis["B9"].protection.locked
    assert all(sheet.protection.sheet for sheet in workbook.worksheets)
    kpis = workbook["Marketing KPIs"]
    assert kpis["B7"].value.startswith("=IFERROR(B5/B6") and "StartDate" in kpis["B5"].value
    assert "'Period analysis'!$B$7" in kpis["B5"].value
    assert len(workbook["Dashboard"]._charts) == 2
    assert "'Period analysis'!" in workbook["Dashboard"]._charts[1].series[0].val.numRef.f
    assert workbook["Dashboard"]["H1"].value.startswith("=MAX(Daily!")
    definitions = [row[0] for row in workbook["Definitions"].iter_rows(min_row=4, values_only=True)]
    assert "Human open rate" in definitions and "Cost per booked call (CPDM)" in definitions
