"""Build the formula-driven Excel workbook from the governed BI CSVs.

Every figure on a report sheet is an Excel formula over the data sheets, written
against named ranges (``pd_spend_cents``, ``WinStart``) so a reviewer can read it.
One control, the reporting window on the Dashboard, drives every windowed figure
in the workbook through the ``Calc`` sheet, which shows how each date is derived.
The Audit sheet holds reconciliation checks that must all equal zero.

The file is built with openpyxl, deterministically, so CI can compare the formulas
of the committed copy with a fresh build. openpyxl cannot calculate, so the
committed copy is then opened and recalculated in Excel once (``--recalculate``,
Windows only) to carry cached values for viewers that do not recalculate.

Regenerate after ``python -m growthops.export_bi --refresh-pbip``.
"""

from __future__ import annotations

import argparse
import csv
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

SOURCE = Path("dashboards/powerbi-data")
OUTPUT = Path("dashboards/GrowthOps_OS_Excel_Dashboard.xlsx")
CATALOG = Path("docs/metric-catalog.md")

# Palette: the report palette's blue and orange, a navy band, quiet greys.
NAVY, INK, MUTED, FAINT = "0F1B2D", "1F2328", "5B6470", "8A929C"
BLUE, ORANGE, GREEN, RED, AMBER = "2F6FD6", "D95926", "1A7F37", "C62828", "B7791F"
TILE, LINE, HEAD, INPUT = "F4F6FA", "D6DBE3", "E8EEF7", "FFF4CC"
FONT = "Segoe UI"

MONEY = '"$"#,##0'
MONEY_2 = '"$"#,##0.00'
COUNT = "#,##0"
PCT = "0.0%"
PCT_2 = "0.00%"
ROAS = '0.00"x"'
DATE = "d mmm yyyy"
# Deltas carry their own colour: green for a move in the good direction, red for the bad one.
UP_GOOD = '[Color10]"▲ "+0.0%;[Red]"▼ "-0.0%;"– "0.0%'
DOWN_GOOD = '[Red]"▲ "+0.0%;[Color10]"▼ "-0.0%;"– "0.0%'
NEUTRAL = '"▲ "+0.0%;"▼ "-0.0%;"– "0.0%'
PTS_UP_GOOD = '[Color10]"▲ "+0.0" pts";[Red]"▼ "-0.0" pts";"– "0.0" pts"'
PTS_DOWN_GOOD = '[Red]"▲ "+0.00" pts";[Color10]"▼ "-0.00" pts";"– "0.00" pts"'

WINDOWS = ("Last 7 days", "Last 28 days", "Last 90 days", "Month to date", "Quarter to date", "Custom")
DEFAULT_WINDOW = "Last 28 days"

# (sheet title, CSV, name prefix for its columns' defined names)
DATA_SHEETS = (
    ("Daily", "mart_growth_daily", "dly"),
    ("Paid daily", "mart_paid_efficiency_daily", "pd"),
    ("Attribution", "fact_cash_attribution", "att"),
    ("Email", "mart_email_performance", "em"),
    ("Campaign dim", "dim_campaign", "dc"),
    ("Campaigns", "mart_campaign_performance", "cp"),
    ("Revenue", "mart_revenue", "rv"),
    ("Bridge", "mart_revenue_bridge", "br"),
    ("Platforms", "mart_platform_comparison", "pl"),
    ("Funnel", "mart_funnel", "fn"),
    ("Experiment", "mart_experiment_variants", "ex"),
    ("Content", "mart_content_performance", "ct"),
    ("Quality", "quality_scorecard", "qc"),
    ("Health", "mart_measurement_health", "mh"),
    ("Migration", "mart_migration_summary", "mg"),
    ("Links", "mart_link_hygiene", "ln"),
    ("Renewals", "mart_renewal_risk", "rn"),
    ("Incidents", "incident_register", "inc"),
)
DATE_COLUMNS = {"day", "due_date", "sent_date", "paid_date", "started_on", "ended_on", "date"}
REPORT_SHEETS = ("Cover", "Dashboard", "Campaign scorecard", "Email health", "Revenue truth", "Funnel and test",
                 "Data quality", "Audit", "Definitions", "Calc")


# --------------------------------------------------------------------------
# Loading and small helpers
# --------------------------------------------------------------------------

def _typed(value: str, column: str):
    if value == "":
        return None
    if column in DATE_COLUMNS:
        return date.fromisoformat(value[:10])
    if value in ("True", "False"):
        return value == "True"
    if column.endswith(("_id", "_label", "_name")) or column in ("title", "subject", "description", "note"):
        return value
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def _load(name: str) -> tuple[list[str], list[list]]:
    with (SOURCE / f"{name}.csv").open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    header = rows[0]
    return header, [[_typed(value, column) for value, column in zip(row, header)] for row in rows[1:]]


def _font(size: float = 10, *, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _rule(color: str = LINE, style: str = "thin") -> Side:
    return Side(style=style, color=color)


class Book:
    """The workbook plus the defined names it has registered."""

    def __init__(self) -> None:
        self.wb = Workbook()
        self.names: dict[str, str] = {}
        self.paid_campaigns = 0

    def name(self, name: str, ref: str) -> None:
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        self.names[name] = ref


def _header_band(sheet, title: str, subtitle: str, width: int) -> None:
    last = get_column_letter(width)
    for row in (1, 2, 3):
        for col in range(1, width + 1):
            sheet.cell(row, col).fill = _fill(NAVY)
    sheet.row_dimensions[1].height = 8
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 20
    sheet["B2"] = title
    sheet["B2"].font = _font(18, bold=True, color="FFFFFF")
    sheet["B2"].alignment = Alignment(vertical="center")
    sheet["B3"] = subtitle
    sheet["B3"].font = _font(9, color="C9D1DC")
    sheet["B3"].alignment = Alignment(vertical="top")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 2
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last}{max(sheet.max_row, 40)}"


def _section(sheet, cell: str, text: str) -> None:
    sheet[cell] = text
    sheet[cell].font = _font(11, bold=True, color=NAVY)


def _table_header(sheet, row: int, col: int, labels: list[str], widths: list[float] | None = None,
                  left: int = 1) -> None:
    """Column headings; the first `left` columns hold text and align left, the rest hold numbers."""
    for offset, label in enumerate(labels):
        cell = sheet.cell(row, col + offset, label)
        cell.font = _font(9, bold=True, color=MUTED)
        cell.fill = _fill(HEAD)
        cell.border = Border(bottom=_rule(FAINT))
        cell.alignment = Alignment(horizontal="left" if offset < left else "right", vertical="center",
                                   wrap_text=True)
        if widths:
            sheet.column_dimensions[get_column_letter(col + offset)].width = widths[offset]
    sheet.row_dimensions[row].height = 32


def _body(cell, fmt: str | None = None, *, bold: bool = False, align: str = "right") -> None:
    cell.font = _font(10, bold=bold)
    cell.border = Border(bottom=_rule())
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        cell.number_format = fmt


def _note(sheet, cell: str, text: str) -> None:
    sheet[cell] = text
    sheet[cell].font = _font(9, italic=True, color=MUTED)


def _window(values: str, dates: str, *, prior: bool = False, extra: str = "") -> str:
    start, end = ("PriorStart", "PriorEnd") if prior else ("WinStart", "WinEnd")
    return f'SUMIFS({values},{dates},">="&{start},{dates},"<="&{end}{extra})'


# --------------------------------------------------------------------------
# Data sheets
# --------------------------------------------------------------------------

def _data_sheet(book: Book, title: str, source: str, prefix: str) -> dict[str, int]:
    sheet = book.wb.create_sheet(title)
    header, rows = _load(source)
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    last = len(rows) + 1
    for index, column in enumerate(header, 1):
        cell = sheet.cell(1, index)
        cell.font = _font(9, bold=True, color="FFFFFF")
        cell.fill = _fill(NAVY)
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(11, min(48, len(column) + 3))
        fmt = None
        if column in DATE_COLUMNS:
            fmt = "yyyy-mm-dd"
        elif column.endswith("_cents"):
            fmt = COUNT
        elif column.endswith("_rate") or column in ("rate", "target"):
            fmt = PCT
        if fmt:
            for row in range(2, last + 1):
                sheet.cell(row, index).number_format = fmt
        # One defined name per column, so report formulas read as words.
        book.name(f"{prefix}_{column}", f"'{title}'!${letter}$2:${letter}${last}")
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.tabColor = "8A929C"
    table = Table(displayName=f"t_{prefix}", ref=f"A1:{get_column_letter(len(header))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    sheet.add_table(table)
    for row in sheet.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = _font(9)
    return {"rows": len(rows), "last": last}


# --------------------------------------------------------------------------
# Calc: the window, and the series the charts plot
# --------------------------------------------------------------------------

def _calc_sheet(book: Book, months: list[date], domains: list[str], bridge: list[list]) -> dict:
    sheet = book.wb["Calc"]
    _header_band(sheet, "Calc", "How the reporting window and chart series are derived. Nothing here is typed.", 12)
    _section(sheet, "B5", "Reporting window")
    rows = [
        ("As-of date (last complete day of cash)", "=MAX(dly_day)", DATE, "AsOf"),
        ("First data day", "=MIN(dly_day)", DATE, "FirstDay"),
        ("Window chosen on the Dashboard", "=Dashboard!$E$5", None, None),
        ("Window end", '=IF($C$8="Custom",Dashboard!$L$5,AsOf)', DATE, "WinEnd"),
        ("Window start",
         '=IF($C$8="Custom",Dashboard!$I$5,CHOOSE(MATCH($C$8,WindowList,0),AsOf-6,AsOf-27,AsOf-89,'
         'DATE(YEAR(AsOf),MONTH(AsOf),1),DATE(YEAR(AsOf),FLOOR(MONTH(AsOf)-1,3)+1,1),AsOf))', DATE, "WinStart"),
        ("Days in window", "=WinEnd-WinStart+1", COUNT, "WinDays"),
        ("Prior window end", "=WinStart-1", DATE, "PriorEnd"),
        ("Prior window start", "=WinStart-WinDays", DATE, "PriorStart"),
        ("Window check",
         '=IF(OR(NOT(ISNUMBER(WinStart)),NOT(ISNUMBER(WinEnd))),"Enter both custom dates",'
         'IF(WinEnd<WinStart,"End is before start",IF(OR(WinStart<FirstDay,WinEnd>AsOf),"Outside the data",'
         'IF(PriorStart<FirstDay,"Ready (no full prior window)","Ready"))))', None, "WindowCheck"),
    ]
    for offset, (label, formula, fmt, name) in enumerate(rows, start=6):
        sheet.cell(offset, 2, label).font = _font(10, color=MUTED)
        cell = sheet.cell(offset, 3, formula)
        _body(cell, fmt, align="left")
        if name:
            book.name(name, f"Calc!$C${offset}")
    book.name("WindowList", f"Calc!$E$6:$E${5 + len(WINDOWS)}")
    sheet["E5"] = "Window options"
    sheet["E5"].font = _font(9, bold=True, color=MUTED)
    for offset, option in enumerate(WINDOWS, start=6):
        sheet.cell(offset, 5, option).font = _font(10)
    # Kept for readers of the previous workbook, whose formulas used these names.
    book.name("StartDate", "Calc!$C$10")
    book.name("EndDate", "Calc!$C$9")

    # Monthly series for the Dashboard chart.
    top = 18
    _section(sheet, f"B{top - 1}", "Monthly series (Dashboard chart)")
    _table_header(sheet, top, 2, ["Month", "Label", "Net cash", "Ad spend", "Paid leads", "Paid MQLs",
                                  "Paid MQL rate"])
    for index, month in enumerate(months):
        row = top + 1 + index
        sheet.cell(row, 2, month).number_format = "yyyy-mm-dd"
        sheet.cell(row, 3, month.strftime("%b %y"))
        nxt = f'"<"&EDATE(B{row},1)'
        sheet.cell(row, 4, f'=SUMIFS(dly_net_cash_cents,dly_day,">="&B{row},dly_day,{nxt})/100').number_format = MONEY
        sheet.cell(row, 5, f'=SUMIFS(dly_spend_cents,dly_day,">="&B{row},dly_day,{nxt})/100').number_format = MONEY
        sheet.cell(row, 6, f'=SUMIFS(pd_leads,pd_day,">="&B{row},pd_day,{nxt})').number_format = COUNT
        sheet.cell(row, 7, f'=SUMIFS(pd_mqls,pd_day,">="&B{row},pd_day,{nxt})').number_format = COUNT
        sheet.cell(row, 8, f'=IFERROR(G{row}/F{row},NA())').number_format = PCT
        for col in range(2, 9):
            sheet.cell(row, col).font = _font(9)
    month_rows = (top + 1, top + len(months))

    # Weekly bounce rate by sending domain, last twelve complete-or-current weeks.
    wtop = month_rows[1] + 4
    _section(sheet, f"B{wtop - 1}", "Weekly bounce rate by sending domain (Email health chart)")
    _table_header(sheet, wtop, 2, ["Week starting", "Label", *domains])
    for index in range(12):
        row = wtop + 1 + index
        sheet.cell(row, 2, f"=AsOf-WEEKDAY(AsOf,3)-7*{11 - index}").number_format = "yyyy-mm-dd"
        sheet.cell(row, 3, f'=TEXT(B{row},"d mmm")')
        for offset, domain in enumerate(domains):
            dates = f'em_sent_date,">="&$B{row},em_sent_date,"<"&$B{row}+7,em_sending_domain,"{domain}"'
            sheet.cell(row, 4 + offset,
                       f"=IFERROR(SUMIFS(em_bounces,{dates})/SUMIFS(em_sends,{dates}),NA())").number_format = PCT_2
        for col in range(2, 4 + len(domains)):
            sheet.cell(row, col).font = _font(9)
    week_rows = (wtop + 1, wtop + 12)

    # Waterfall helper: an invisible base plus up, down and total bars.
    btop = week_rows[1] + 4
    _section(sheet, f"B{btop - 1}", "Revenue bridge helper (waterfall on Revenue truth)")
    _table_header(sheet, btop, 2, ["Step", "Amount", "Running end", "Base", "Increase", "Decrease", "Total"])
    for index, _row in enumerate(bridge):
        row = btop + 1 + index
        source = index + 2
        sheet.cell(row, 2, f"=Bridge!E{source}")
        sheet.cell(row, 3, f"=Bridge!D{source}/100").number_format = MONEY
        is_total = f'Bridge!C{source}="total"'
        previous = "0" if index == 0 else f"D{row - 1}"
        sheet.cell(row, 4, f"=IF({is_total},C{row},{previous}+C{row})").number_format = MONEY
        sheet.cell(row, 5, f"=IF({is_total},0,MIN({previous},D{row}))").number_format = MONEY
        sheet.cell(row, 6, f"=IF({is_total},0,MAX(C{row},0))").number_format = MONEY
        sheet.cell(row, 7, f"=IF({is_total},0,MAX(-C{row},0))").number_format = MONEY
        sheet.cell(row, 8, f"=IF({is_total},C{row},0)").number_format = MONEY
        for col in range(2, 9):
            sheet.cell(row, col).font = _font(9)
    bridge_rows = (btop + 1, btop + len(bridge))
    for col, width in zip("BCDEFGH", (40, 18, 18, 14, 14, 14, 14)):
        sheet.column_dimensions[col].width = width
    sheet.sheet_properties.tabColor = "8A929C"
    return {"months": month_rows, "weeks": week_rows, "bridge": bridge_rows, "domains": domains}


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------

TILE_COLUMNS = ("B", "E", "H", "K", "N", "Q")
PLATFORM_TOP = 18  # where the platform table starts on Revenue truth; the Dashboard ROAS chart reads it


def _tile(sheet, col: str, label: str, current: str, prior: str, value_fmt: str, delta: str, delta_fmt: str) -> None:
    left = sheet[f"{col}8"].column
    right = get_column_letter(left + 1)
    span = lambda r: f"{col}{r}:{right}{r}"  # noqa: E731
    for r in range(8, 13):
        sheet.merge_cells(span(r))
        for c in (left, left + 1):
            cell = sheet.cell(r, c)
            cell.fill = _fill(TILE)
            cell.border = Border(left=_rule(BLUE, "thick") if c == left else None)
    sheet[f"{col}8"] = label
    sheet[f"{col}8"].font = _font(9, color=MUTED)
    sheet[f"{col}9"] = current
    sheet[f"{col}9"].font = _font(20, bold=True, color=NAVY)
    sheet[f"{col}9"].number_format = value_fmt
    sheet[f"{col}10"] = "=" + f'"prior window: "&TEXT({prior[1:]},"{_text_format(value_fmt)}")'
    sheet[f"{col}10"].font = _font(8, color=FAINT)
    sheet[f"{col}11"] = delta
    sheet[f"{col}11"].font = _font(10, bold=True)
    sheet[f"{col}11"].number_format = delta_fmt
    sheet[f"{col}12"] = "vs the prior window"
    sheet[f"{col}12"].font = _font(8, color=FAINT)
    for r in range(8, 13):
        sheet[f"{col}{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)


def _text_format(fmt: str) -> str:
    return fmt.replace('"', '""')


def _dashboard(book: Book, calc: dict) -> None:
    sheet = book.wb["Dashboard"]
    width = 19
    for c in range(2, width + 1):
        sheet.column_dimensions[get_column_letter(c)].width = 10.5
    for c in ("D", "G", "J", "M", "P"):
        sheet.column_dimensions[c].width = 2.5
    _header_band(sheet, "ScaleLab GrowthOps · Executive dashboard",
                 "Synthetic portfolio case · USD · every figure is a formula over the dbt marts · change the "
                 "window below and every report sheet follows", width)
    sheet["Q2"] = '="Data through "&TEXT(AsOf,"d mmm yyyy")'
    sheet["Q2"].font = _font(10, bold=True, color="FFFFFF")
    sheet["Q2"].alignment = Alignment(horizontal="right", vertical="center")
    sheet.merge_cells("Q2:S2")

    # Controls.
    sheet["B5"] = "Reporting window"
    sheet["B5"].font = _font(10, bold=True)
    sheet.merge_cells("B5:C5")
    sheet["E5"] = DEFAULT_WINDOW
    sheet.merge_cells("E5:F5")
    sheet["H5"] = "Custom start"
    sheet["I5"] = date(2026, 9, 1)
    sheet["K5"] = "Custom end"
    sheet["L5"] = date(2026, 9, 25)
    sheet.merge_cells("I5:J5")
    sheet.merge_cells("L5:M5")
    for cell in ("E5", "I5", "L5"):
        sheet[cell].fill = _fill(INPUT)
        sheet[cell].font = _font(10, bold=True)
        sheet[cell].border = Border(left=_rule(AMBER), right=_rule(AMBER), top=_rule(AMBER), bottom=_rule(AMBER))
        sheet[cell].alignment = Alignment(horizontal="center")
        sheet[cell].protection = Protection(locked=False)
    sheet["F5"].fill = _fill(INPUT)
    sheet["I5"].number_format = sheet["L5"].number_format = DATE
    for cell in ("H5", "K5"):
        sheet[cell].font = _font(9, color=MUTED)
        sheet[cell].alignment = Alignment(horizontal="right")
    sheet["N5"] = ('=TEXT(WinStart,"d mmm yyyy")&" – "&TEXT(WinEnd,"d mmm yyyy")&"  ("&WinDays&" days) · prior: "'
                   '&TEXT(PriorStart,"d mmm")&" – "&TEXT(PriorEnd,"d mmm")')
    sheet["N5"].font = _font(9, color=MUTED)
    sheet.merge_cells("N5:S5")
    sheet["B6"] = '=IF(LEFT(WindowCheck,5)="Ready","","⚠ "&WindowCheck)'
    sheet["B6"].font = _font(9, bold=True, color=RED)
    choice = DataValidation(type="list", formula1="=WindowList", showErrorMessage=True,
                            errorTitle="Reporting window", error="Choose a window from the list.")
    sheet.add_data_validation(choice)
    choice.add("E5")
    for address in ("I5", "L5"):
        dates = DataValidation(type="date", operator="between", formula1="FirstDay", formula2="AsOf",
                               showErrorMessage=True, errorTitle="Custom dates",
                               error="Choose a date inside the data (see Calc for the range).")
        sheet.add_data_validation(dates)
        dates.add(address)

    def paid(column: str, prior: bool = False) -> str:
        return _window(f"pd_{column}", "pd_day", prior=prior)

    def daily(column: str, prior: bool = False) -> str:
        return _window(f"dly_{column}", "dly_day", prior=prior)

    def email(column: str, prior: bool = False) -> str:
        return _window(f"em_{column}", "em_sent_date", prior=prior)

    tiles = [
        ("Net cash", f"={daily('net_cash_cents')}/100", f"={daily('net_cash_cents', True)}/100", MONEY, "pct", UP_GOOD),
        ("Paid spend", f"={paid('spend_cents')}/100", f"={paid('spend_cents', True)}/100", MONEY, "pct", NEUTRAL),
        ("Cost per paid lead", f"=IFERROR({paid('spend_cents')}/100/{paid('leads')},NA())",
         f"=IFERROR({paid('spend_cents', True)}/100/{paid('leads', True)},NA())", MONEY_2, "pct", DOWN_GOOD),
        ("Paid MQL rate", f"=IFERROR({paid('mqls')}/{paid('leads')},NA())",
         f"=IFERROR({paid('mqls', True)}/{paid('leads', True)},NA())", PCT, "pts", PTS_UP_GOOD),
        ("Cost per booked call", f"=IFERROR({paid('spend_cents')}/100/{paid('calls_booked')},NA())",
         f"=IFERROR({paid('spend_cents', True)}/100/{paid('calls_booked', True)},NA())", MONEY, "pct", DOWN_GOOD),
        ("Email bounce rate", f"=IFERROR({email('bounces')}/{email('sends')},NA())",
         f"=IFERROR({email('bounces', True)}/{email('sends', True)},NA())", PCT_2, "pts", PTS_DOWN_GOOD),
    ]
    calc_sheet = book.wb["Calc"]
    _section(calc_sheet, "J5", "Prior-window values for the Dashboard tiles")
    for offset, (label, current, prior, fmt, kind, delta_fmt) in enumerate(tiles):
        # The prior value lives on Calc so the tile can show it and the delta can use it.
        row = 6 + offset
        calc_sheet.cell(row, 10, label).font = _font(10, color=MUTED)
        _body(calc_sheet.cell(row, 11, prior), fmt, align="left")
        col = TILE_COLUMNS[offset]
        if kind == "pct":
            delta = f"=IFERROR({col}9/Calc!$K${row}-1,NA())"
        else:
            delta = f"=IFERROR(({col}9-Calc!$K${row})*100,NA())"
        _tile(sheet, col, label, current, f"=Calc!$K${row}", fmt, delta, delta_fmt)
    calc_sheet.column_dimensions["J"].width = 24
    calc_sheet.column_dimensions["K"].width = 14

    # What changed, as sentences written by formulas.
    _section(sheet, "B14", "What changed in the window")
    names = f"'Campaign scorecard'!$B$7:$B${6 + book.paid_campaigns}"
    rates = f"'Campaign scorecard'!$Q$7:$Q${6 + book.paid_campaigns}"
    # MIN ignores the text left by campaigns under 50 leads. (AGGREGATE would need
    # the _xlfn. prefix openpyxl does not write, and reads #NAME? in Excel.)
    worst = (f'IFERROR(INDEX({names},MATCH(MIN({rates}),{rates},0))&" at "&TEXT(MIN({rates}),"0.0%"),'
             '"none with 50+ leads")')
    sentences = [
        '="Net cash was "&TEXT(B9,"$#,##0")&" ("&TEXT(B11,"+0%;-0%")&" on the prior window) on "'
        '&TEXT(E9,"$#,##0")&" of paid spend."',
        '="Paid media bought leads at "&TEXT(H9,"$#,##0.00")&" each, and "&TEXT(K9,"0.0%")&" of them '
        'qualified ("&TEXT(Calc!$K$9,"0.0%")&" in the prior window). Lowest MQL rate at scale '
        '(50+ leads): "&' + worst + '&"."',
        '="Email bounced at "&TEXT(Q9,"0.00%")&IF(Q9>0.02,", above the 2% limit: check the sending domain on '
        'Email health.",", inside the 2% limit.")',
    ]
    for offset, formula in enumerate(sentences, start=15):
        sheet.merge_cells(f"B{offset}:S{offset}")
        sheet[f"B{offset}"] = formula
        sheet[f"B{offset}"].font = _font(10)
        sheet[f"B{offset}"].alignment = Alignment(vertical="center", indent=1)
        sheet.row_dimensions[offset].height = 20

    # Charts.
    first, last = calc["months"]
    monthly = BarChart()
    monthly.type = "col"
    monthly.grouping = "clustered"
    monthly.title = "Net cash and ad spend by month (bars) · paid MQL rate (line)"
    monthly.add_data(Reference(book.wb["Calc"], min_col=4, max_col=5, min_row=first - 1, max_row=last),
                     titles_from_data=True)
    monthly.set_categories(Reference(book.wb["Calc"], min_col=3, min_row=first, max_row=last))
    _style_series(monthly, (BLUE, ORANGE))
    rate = LineChart()
    rate.add_data(Reference(book.wb["Calc"], min_col=8, min_row=first - 1, max_row=last), titles_from_data=True)
    rate.y_axis.axId = 200
    rate.y_axis.number_format = "0%"
    rate.y_axis.crosses = "max"
    rate.y_axis.delete = False
    rate.y_axis.majorGridlines = None
    rate.y_axis.title = None
    rate.series[0].graphicalProperties.line.solidFill = GREEN
    rate.series[0].graphicalProperties.line.width = 28000
    rate.series[0].smooth = False
    monthly.y_axis.number_format = '"$"#,##0,"K"'
    monthly += rate
    _chart_frame(monthly, 17.5, 8.5)
    sheet.add_chart(monthly, "B19")

    roas = BarChart()
    roas.type = "bar"
    roas.title = "ROAS by platform: what the platform reports vs net cash"
    head = PLATFORM_TOP + 1
    roas.add_data(Reference(book.wb["Revenue truth"], min_col=7, max_col=8, min_row=head, max_row=head + 3),
                  titles_from_data=True)
    roas.set_categories(Reference(book.wb["Revenue truth"], min_col=2, min_row=head + 1, max_row=head + 3))
    _style_series(roas, (ORANGE, BLUE))
    roas.y_axis.number_format = '0.0"x"'
    _value_labels(roas, '0.0"x"')
    _chart_frame(roas, 12.5, 8.5)
    sheet.add_chart(roas, "M19")
    sheet.freeze_panes = "A7"
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:S37"
    sheet.sheet_properties.tabColor = BLUE


def _value_labels(chart, fmt: str) -> None:
    """Value only. Left unset, Excel also prints the series and category name on every bar."""
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
        setattr(chart.dataLabels, flag, False)
    chart.dataLabels.numFmt = fmt


def _style_series(chart, colours) -> None:
    for series, colour in zip(chart.series, colours):
        series.graphicalProperties = GraphicalProperties(solidFill=colour)
        series.graphicalProperties.line.solidFill = colour


def _chart_frame(chart, width: float, height: float) -> None:
    chart.width, chart.height = width, height
    chart.legend.position = "b"
    chart.style = 2
    if chart.y_axis.majorGridlines is not None:
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E9EF"))
    chart.x_axis.delete = False
    chart.y_axis.delete = False


# --------------------------------------------------------------------------
# Campaign scorecard
# --------------------------------------------------------------------------

def _scorecard(book: Book, paid_campaigns: list[tuple[str, str]]) -> None:
    sheet = book.wb["Campaign scorecard"]
    _header_band(sheet, "Campaign scorecard",
                 '=\"Paid campaigns in the window \"&TEXT(WinStart,\"d mmm\")&\" – \"&TEXT(WinEnd,\"d mmm yyyy\")'
                 '&\" · activity basis: each lead, MQL and call is credited to the campaign that created the lead\"',
                 17)
    labels = ["Campaign", "Platform", "Spend", "Impressions", "Clicks", "CTR", "Leads", "Cost per lead",
              "MQLs", "MQL rate", "Calls booked", "Cost per call", "Net cash (payment date)", "Cash ROAS",
              "Flag"]
    widths = [26, 10, 12, 13, 10, 8, 9, 11, 9, 9, 9, 11, 14, 10, 24]
    _table_header(sheet, 6, 2, labels, widths, left=2)
    sheet.cell(6, 16).alignment = Alignment(horizontal="left", vertical="center")
    first = 7
    last = first + len(paid_campaigns) - 1
    for offset, (campaign, platform) in enumerate(paid_campaigns):
        row = first + offset
        by = f',pd_campaign_id,$B{row}'
        sheet.cell(row, 2, campaign)
        sheet.cell(row, 3, platform)
        cells = [
            (4, f"={_window('pd_spend_cents', 'pd_day', extra=by)}/100", MONEY),
            (5, f"={_window('pd_impressions', 'pd_day', extra=by)}", COUNT),
            (6, f"={_window('pd_clicks', 'pd_day', extra=by)}", COUNT),
            (7, f"=IFERROR(F{row}/E{row},\"–\")", PCT_2),
            (8, f"={_window('pd_leads', 'pd_day', extra=by)}", COUNT),
            (9, f"=IFERROR(D{row}/H{row},\"–\")", MONEY_2),
            (10, f"={_window('pd_mqls', 'pd_day', extra=by)}", COUNT),
            (11, f"=IFERROR(J{row}/H{row},\"–\")", PCT),
            (12, f"={_window('pd_calls_booked', 'pd_day', extra=by)}", COUNT),
            (13, f"=IFERROR(D{row}/L{row},\"–\")", MONEY),
            (14, f"={_window('att_net_cash_cents', 'att_paid_date', extra=f',att_campaign_id,$B{row}')}/100", MONEY),
            (15, f"=IFERROR(N{row}/D{row},\"–\")", ROAS),
            (16, f'=IF(D{row}=0,"No spend",IF(IFERROR(K{row}<0.75*$K${last + 1},FALSE),"Low lead quality",'
                 f'IF(IFERROR(O{row}<1,FALSE),"Returned less than it cost","")))', None),
        ]
        # Hidden helper: the MQL rate of campaigns with 50+ leads, so "lowest quality" means at scale.
        cells.append((17, f'=IF(AND(ISNUMBER(K{row}),H{row}>=50),K{row},"")', PCT))
        for col, formula, fmt in cells:
            sheet.cell(row, col, formula)
        for col in range(2, 18):
            _body(sheet.cell(row, col), next((f for c, _, f in cells if c == col), None),
                  align="left" if col in (2, 3, 16) else "right")
        sheet.cell(row, 16).font = _font(9, bold=True, color=RED)
    total = last + 1
    sheet.cell(total, 2, "All paid campaigns")
    for col, letter in ((4, "D"), (5, "E"), (6, "F"), (8, "H"), (10, "J"), (12, "L"), (14, "N")):
        sheet.cell(total, col, f"=SUM({letter}{first}:{letter}{last})")
    for col, formula in ((7, f"=IFERROR(F{total}/E{total},NA())"), (9, f"=IFERROR(D{total}/H{total},NA())"),
                         (11, f"=IFERROR(J{total}/H{total},NA())"), (13, f"=IFERROR(D{total}/L{total},NA())"),
                         (15, f"=IFERROR(N{total}/D{total},NA())")):
        sheet.cell(total, col, formula)
    formats = {4: MONEY, 5: COUNT, 6: COUNT, 7: PCT_2, 8: COUNT, 9: MONEY_2, 10: COUNT, 11: PCT, 12: COUNT,
               13: MONEY, 14: MONEY, 15: ROAS}
    for col in range(2, 17):
        cell = sheet.cell(total, col)
        _body(cell, formats.get(col), bold=True, align="left" if col in (2, 3) else "right")
        cell.fill = _fill(HEAD)
        cell.border = Border(top=_rule(NAVY), bottom=_rule(NAVY))
    span = f"{first}:{last}"
    sheet.conditional_formatting.add(f"D{first}:D{last}", DataBarRule(start_type="num", start_value=0,
                                                                       end_type="max", color="9DBBEB"))
    sheet.conditional_formatting.add(f"I{first}:I{last}", ColorScaleRule(
        start_type="min", start_color="D8F0DE", mid_type="percentile", mid_value=50, mid_color="FFFFFF",
        end_type="max", end_color="F8D7D5"))
    sheet.conditional_formatting.add(f"K{first}:K{last}", ColorScaleRule(
        start_type="min", start_color="F8D7D5", mid_type="percentile", mid_value=50, mid_color="FFFFFF",
        end_type="max", end_color="D8F0DE"))
    sheet.conditional_formatting.add(f"O{first}:O{last}", CellIsRule(operator="lessThan", formula=["1"],
                                                                     font=Font(name=FONT, color=RED, bold=True)))
    _note(sheet, f"B{total + 2}",
          "Cost per lead: green is cheaper. MQL rate: green qualifies more. Flags: MQL rate under three-quarters "
          "of the paid average, or cash ROAS under 1.0x.")
    _note(sheet, f"B{total + 3}",
          "Cash is credited to the campaign that created the buyer's lead, on the payment date, so a short window "
          "understates ROAS for campaigns whose buyers pay later. Attribution is descriptive, not incremental.")
    del span

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Cost per booked call by campaign (window)"
    chart.add_data(Reference(sheet, min_col=13, min_row=6, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=2, min_row=first, max_row=last))
    _style_series(chart, (BLUE,))
    chart.legend = None
    _value_labels(chart, '"$"#,##0')
    chart.y_axis.number_format = '"$"#,##0'
    _chart_frame_nolegend(chart, 16, 7.5)
    sheet.add_chart(chart, f"B{total + 5}")
    sheet.column_dimensions["Q"].hidden = True
    sheet.freeze_panes = "C7"
    sheet.sheet_properties.tabColor = BLUE


def _chart_frame_nolegend(chart, width: float, height: float) -> None:
    chart.width, chart.height = width, height
    chart.style = 2
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    if chart.y_axis.majorGridlines is not None:
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E9EF"))


# --------------------------------------------------------------------------
# Email health
# --------------------------------------------------------------------------

def _email(book: Book, calc: dict, types: list[str]) -> None:
    sheet = book.wb["Email health"]
    domains = calc["domains"]
    _header_band(sheet, "Email health",
                 '="Sends in the window "&TEXT(WinStart,"d mmm")&" – "&TEXT(WinEnd,"d mmm yyyy")&'
                 '" · human opens exclude privacy-proxy machine opens · limits: bounce 2%, complaints 0.1%"', 12)
    labels = ["", "Sent", "Delivered", "Bounce rate", "Human open rate", "Reported open rate", "Click rate",
              "Click-to-open", "Complaint rate", "Unsubscribe rate", "Status"]
    widths = [26, 11, 11, 11, 12, 12, 10, 11, 12, 12, 20]

    def block(top: int, heading: str, field: str, members: list[str]) -> int:
        _section(sheet, f"B{top}", heading)
        labels[0] = heading.split(" by ")[-1].capitalize()
        _table_header(sheet, top + 1, 2, labels, widths)
        sheet.cell(top + 1, 12).alignment = Alignment(horizontal="left", vertical="center")
        for offset, member in enumerate(members):
            row = top + 2 + offset
            by = f',em_{field},$B{row}'
            s = lambda column: _window(f"em_{column}", "em_sent_date", extra=by)  # noqa: E731
            sheet.cell(row, 2, member)
            cells = [
                (3, f"={s('sends')}", COUNT), (4, f"={s('delivered')}", COUNT),
                (5, f"=IFERROR({s('bounces')}/C{row},NA())", PCT_2),
                (6, f"=IFERROR({s('human_opens')}/D{row},NA())", PCT),
                (7, f"=IFERROR({s('opens')}/D{row},NA())", PCT),
                (8, f"=IFERROR({s('clicks')}/D{row},NA())", PCT_2),
                (9, f"=IFERROR({s('clicks')}/{s('human_opens')},NA())", PCT),
                (10, f"=IFERROR({s('spam_complaints')}/D{row},NA())", "0.000%"),
                (11, f"=IFERROR({s('unsubscribes')}/D{row},NA())", PCT_2),
                (12, f'=IF(C{row}=0,"No sends",IF(OR(IFERROR(E{row}>0.02,FALSE),IFERROR(J{row}>0.001,FALSE)),'
                     f'"Over limit","OK"))', None),
            ]
            for col, formula, fmt in cells:
                sheet.cell(row, col, formula)
                _body(sheet.cell(row, col), fmt, align="left" if col == 12 else "right")
            _body(sheet.cell(row, 2), align="left")
        first, last = top + 2, top + 1 + len(members)
        sheet.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(
            operator="greaterThan", formula=["0.02"], font=Font(name=FONT, color=RED, bold=True),
            fill=_fill("FBE3E1")))
        sheet.conditional_formatting.add(f"J{first}:J{last}", CellIsRule(
            operator="greaterThan", formula=["0.001"], font=Font(name=FONT, color=RED, bold=True),
            fill=_fill("FBE3E1")))
        sheet.conditional_formatting.add(f"L{first}:L{last}", FormulaRule(
            formula=[f'L{first}="Over limit"'], font=Font(name=FONT, color=RED, bold=True)))
        return last

    end = block(5, "By sending domain", "sending_domain", domains)
    end = block(end + 3, "By email type", "email_type", types)

    first, last = calc["weeks"]
    chart = LineChart()
    chart.title = "Weekly bounce rate by sending domain (last 12 weeks; limit 2%)"
    chart.add_data(Reference(book.wb["Calc"], min_col=4, max_col=3 + len(domains), min_row=first - 1,
                             max_row=last), titles_from_data=True)
    chart.set_categories(Reference(book.wb["Calc"], min_col=3, min_row=first, max_row=last))
    for series, colour in zip(chart.series, (BLUE, RED, ORANGE)):
        series.graphicalProperties.line.solidFill = colour
        series.graphicalProperties.line.width = 28000
        series.smooth = False
    chart.y_axis.number_format = "0.0%"
    _chart_frame(chart, 18, 7.5)
    sheet.add_chart(chart, f"B{end + 3}")
    sheet.sheet_properties.tabColor = BLUE


# --------------------------------------------------------------------------
# Revenue truth
# --------------------------------------------------------------------------

def _revenue(book: Book, calc: dict, bridge: list[list], platforms: int) -> None:
    sheet = book.wb["Revenue truth"]
    _header_band(sheet, "Revenue truth",
                 "Which revenue number is right? All time. Platforms claim, the CRM books, the bank collects; "
                 "every step between them is computed and the residual is zero.", 12)
    _section(sheet, "B5", "From CRM bookings to net collected cash")
    _table_header(sheet, 6, 2, ["Step", "Amount", "What it is"])
    sheet.cell(6, 4).alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells("D6:I6")
    for index in range(len(bridge)):
        row = 7 + index
        source = index + 2
        total = bridge[index][2] == "total"
        sheet.cell(row, 2, f"=Bridge!E{source}")
        sheet.cell(row, 3, f"=Bridge!D{source}/100")
        sheet.cell(row, 4, f"=Bridge!F{source}")
        sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=9)
        _body(sheet.cell(row, 2), align="left", bold=total)
        _body(sheet.cell(row, 3), MONEY, bold=total)
        _body(sheet.cell(row, 4), align="left")
        if total:
            for col in (2, 3, 4):
                sheet.cell(row, col).fill = _fill(HEAD)
    after = 7 + len(bridge)
    sheet.cell(after, 2, "Residual (must be $0)")
    sheet.cell(after, 3, "=Audit!C8+Audit!C9")
    _body(sheet.cell(after, 2), align="left")
    _body(sheet.cell(after, 3), MONEY, bold=True)

    top = PLATFORM_TOP
    _section(sheet, f"B{top}", "What each platform claims against the cash it bought")
    _table_header(sheet, top + 1, 2, ["Platform", "Spend", "Platform-reported value", "Warehouse net cash",
                                      "Overstatement", "Reported ROAS", "Warehouse ROAS", "Claim multiple"],
                  [24, 13, 16, 16, 14, 12, 12, 12])
    for index in range(platforms):
        row = top + 2 + index
        source = index + 2
        cells = [(2, f"=Platforms!A{source}", None), (3, f"=Platforms!B{source}/100", MONEY),
                 (4, f"=Platforms!D{source}/100", MONEY), (5, f"=Platforms!E{source}/100", MONEY),
                 (6, f"=D{row}-E{row}", MONEY), (7, f"=IFERROR(D{row}/C{row},NA())", ROAS),
                 (8, f"=IFERROR(E{row}/C{row},NA())", ROAS), (9, f"=IFERROR(D{row}/E{row},NA())", ROAS)]
        for col, formula, fmt in cells:
            sheet.cell(row, col, formula)
            _body(sheet.cell(row, col), fmt, align="left" if col == 2 else "right")
    total = top + 2 + platforms
    sheet.cell(total, 2, "All paid platforms")
    for col, letter in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
        sheet.cell(total, col, f"=SUM({letter}{top + 2}:{letter}{total - 1})")
    for col, formula in ((7, f"=D{total}/C{total}"), (8, f"=E{total}/C{total}"), (9, f"=D{total}/E{total}")):
        sheet.cell(total, col, formula)
    for col in range(2, 10):
        _body(sheet.cell(total, col), {3: MONEY, 4: MONEY, 5: MONEY, 6: MONEY, 7: ROAS, 8: ROAS, 9: ROAS}.get(col),
              bold=True, align="left" if col == 2 else "right")
        sheet.cell(total, col).fill = _fill(HEAD)
    _note(sheet, f"B{total + 1}",
          "Platforms credit themselves with any conversion they touched; the warehouse credits each payment "
          "once, to the campaign that created the lead.")

    first, last = calc["bridge"]
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "CRM bookings to net cash (axis from $2.8M)"
    chart.add_data(Reference(book.wb["Calc"], min_col=5, max_col=8, min_row=first - 1, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(book.wb["Calc"], min_col=2, min_row=first, max_row=last))
    base, up, down, tot = chart.series
    base.graphicalProperties = GraphicalProperties(noFill=True)
    base.graphicalProperties.line.noFill = True
    up.graphicalProperties = GraphicalProperties(solidFill=GREEN)
    down.graphicalProperties = GraphicalProperties(solidFill=RED)
    tot.graphicalProperties = GraphicalProperties(solidFill=BLUE)
    chart.y_axis.scaling.min = 2_800_000
    chart.y_axis.number_format = '"$"#,##0.0,,"M"'
    chart.gapWidth = 60
    _chart_frame(chart, 18, 8)
    chart.legend = None
    sheet.add_chart(chart, f"B{total + 3}")
    sheet.sheet_properties.tabColor = BLUE


# --------------------------------------------------------------------------
# Funnel and test
# --------------------------------------------------------------------------

def _funnel(book: Book, stages: int) -> None:
    sheet = book.wb["Funnel and test"]
    _header_band(sheet, "Funnel and CTA test",
                 "All time. People reaching each lifecycle stage, and the visitor-randomized CTA test decided on "
                 "cash per visitor rather than on leads.", 15)
    _section(sheet, "B5", "Lead to renewal")
    _table_header(sheet, 6, 2, ["Stage", "People", "From previous stage", "Of all leads"], [22, 12, 22, 13])
    for index in range(stages):
        row = 7 + index
        source = index + 2
        cells = [(2, f"=Funnel!E{source}", None), (3, f"=Funnel!C{source}", COUNT),
                 (4, "" if index == 0 else f"=IFERROR(C{row}/C{row - 1},NA())", PCT),
                 (5, f"=IFERROR(C{row}/$C$7,NA())", PCT_2)]
        for col, formula, fmt in cells:
            if formula:
                sheet.cell(row, col, formula)
            _body(sheet.cell(row, col), fmt, align="left" if col == 2 else "right")
    last = 6 + stages
    sheet.conditional_formatting.add(f"C7:C{last}", DataBarRule(start_type="num", start_value=0,
                                                                  end_type="max", color="9DBBEB"))
    chart = BarChart()
    chart.type = "bar"
    chart.title = "People at each stage"
    chart.add_data(Reference(sheet, min_col=3, min_row=6, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=2, min_row=7, max_row=last))
    chart.x_axis.scaling.orientation = "maxMin"
    _style_series(chart, (BLUE,))
    chart.legend = None
    _value_labels(chart, "#,##0")
    chart.y_axis.crosses = "max"  # reversed categories move the value axis to the top otherwise
    _chart_frame_nolegend(chart, 13, 7.5)
    sheet.add_chart(chart, "G5")

    top = last + 3
    _section(sheet, f"B{top}", "CTA test: control against variant")
    _table_header(sheet, top + 1, 2, ["Variant", "Visitors", "Leads", "Lead rate", "MQLs per lead", "Customers",
                                      "Cash per visitor"], [30, 12, 10, 12, 12, 12, 14])
    for index in range(2):
        row = top + 2 + index
        source = index + 2
        cells = [(2, f"=Experiment!C{source}", None), (3, f"=Experiment!E{source}", COUNT),
                 (4, f"=Experiment!F{source}", COUNT), (5, f"=IFERROR(D{row}/C{row},NA())", PCT_2),
                 (6, f"=IFERROR(Experiment!G{source}/D{row},NA())", PCT), (7, f"=Experiment!H{source}", COUNT),
                 (8, f"=IFERROR(Experiment!I{source}/100/C{row},NA())", MONEY_2)]
        for col, formula, fmt in cells:
            sheet.cell(row, col, formula)
            _body(sheet.cell(row, col), fmt, align="left" if col == 2 else "right")
    lift = top + 4
    sheet.cell(lift, 2, "Variant against control")
    sheet.cell(lift, 5, f"=IFERROR(E{lift - 1}/E{lift - 2}-1,NA())")
    sheet.cell(lift, 8, f"=IFERROR(H{lift - 1}/H{lift - 2}-1,NA())")
    for col in range(2, 9):
        _body(sheet.cell(lift, col), NEUTRAL if col in (5, 8) else None, bold=True,
              align="left" if col == 2 else "right")
        sheet.cell(lift, col).fill = _fill(HEAD)
    sheet.cell(lift + 1, 2,
               f'="Decision: "&IF(AND(E{lift}>0,H{lift}<=0),"keep the control. The variant lifts leads "'
               f'&TEXT(E{lift},"0%")&" but cash per visitor moves "&TEXT(H{lift},"+0%;-0%")&" on only "'
               f'&G{lift - 2}+G{lift - 1}&" buyers.","review with the API\'s bootstrap interval before shipping.")')
    sheet.cell(lift + 1, 2).font = _font(10, bold=True, color=NAVY)
    sheet.sheet_properties.tabColor = BLUE


# --------------------------------------------------------------------------
# Data quality
# --------------------------------------------------------------------------

def _quality(book: Book, checks: int, links: int, incidents: int) -> None:
    sheet = book.wb["Data quality"]
    _header_band(sheet, "Data quality and operations",
                 "Tracking, CRM and migration checks against their targets; the short-link registry; the incident "
                 "register; the renewal queue.", 12)
    _section(sheet, "B5", "Checks against target")
    _table_header(sheet, 6, 2, ["Area", "Check", "Pass rate", "Target", "Gap", "Status"], [14, 26, 11, 10, 10, 16],
                  left=2)
    for index in range(checks):
        row = 7 + index
        source = index + 2
        cells = [(2, f"=Quality!B{source}", None), (3, f"=Quality!C{source}", None),
                 (4, f"=Quality!D{source}", PCT), (5, f"=Quality!E{source}", PCT),
                 (6, f"=(D{row}-E{row})*100", '+0.0" pts";-0.0" pts"'),
                 (7, f'=IF(D{row}>=E{row},"✓ On target","✗ Below target")', None)]
        for col, formula, fmt in cells:
            sheet.cell(row, col, formula)
            _body(sheet.cell(row, col), fmt, align="left" if col in (2, 3, 7) else "right")
    last = 6 + checks
    sheet.conditional_formatting.add(f"G7:G{last}", FormulaRule(formula=['LEFT(G7,1)="✗"'],
                                                                  font=Font(name=FONT, color=RED, bold=True)))
    sheet.conditional_formatting.add(f"G7:G{last}", FormulaRule(formula=['LEFT(G7,1)="✓"'],
                                                                  font=Font(name=FONT, color=GREEN, bold=True)))
    sheet.cell(last + 1, 2, "Checks below target")
    sheet.cell(last + 1, 7, f'=COUNTIF(G7:G{last},"✗*")&" of {checks}"')
    for col in range(2, 8):
        _body(sheet.cell(last + 1, col), bold=True, align="left" if col < 7 else "right")
        sheet.cell(last + 1, col).fill = _fill(HEAD)

    top = last + 4
    _section(sheet, f"B{top}", "Short links against the campaign registry")
    _table_header(sheet, top + 1, 2, ["Link", "Channel", "UTM campaign", "Missing UTM", "Unregistered",
                                      "Off taxonomy", "Clicks, last 30 days", "Defect"],
                  [14, 26, 22, 11, 10, 14, 14, 12], left=3)
    for index in range(links):
        row = top + 2 + index
        source = index + 2
        cells = [(2, f"=Links!A{source}", None), (3, f"=Links!B{source}", None),
                 (4, f'=IF(Links!E{source}="","(none)",Links!E{source})', None),
                 (5, f"=Links!F{source}", None), (6, f"=Links!G{source}", None), (7, f"=Links!H{source}", None),
                 (8, f"=Links!J{source}", COUNT), (9, f'=IF(OR(E{row},F{row},G{row}),"Fix","")', None)]
        for col, formula, fmt in cells:
            sheet.cell(row, col, formula)
            _body(sheet.cell(row, col), fmt, align="left" if col in (2, 3, 4, 9) else "right")
        sheet.cell(row, 9).font = _font(10, bold=True, color=RED)
    llast = top + 1 + links
    sheet.cell(llast + 1, 2, "Share of recent clicks on defective links")
    sheet.cell(llast + 1, 8, f'=IFERROR(SUMIF(I{top + 2}:I{llast},"Fix",H{top + 2}:H{llast})/SUM(H{top + 2}:H{llast}),NA())')
    for col in range(2, 10):
        _body(sheet.cell(llast + 1, col), PCT if col == 8 else None, bold=True, align="left" if col < 8 else "right")
        sheet.cell(llast + 1, col).fill = _fill(HEAD)

    top = llast + 4
    _section(sheet, f"B{top}", "Incident register")
    _table_header(sheet, top + 1, 2, ["Started", "Ended", "Kind", "Description"])
    sheet.merge_cells(start_row=top + 1, start_column=5, end_row=top + 1, end_column=12)
    for index in range(incidents):
        row = top + 2 + index
        source = index + 2
        sheet.cell(row, 2, f"=Incidents!C{source}")
        sheet.cell(row, 3, f'=IF(Incidents!D{source}="","open",Incidents!D{source})')
        sheet.cell(row, 4, f"=Incidents!B{source}")
        sheet.cell(row, 5, f"=Incidents!G{source}")
        sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=12)
        for col in (2, 3):
            _body(sheet.cell(row, col), DATE, align="left")
        _body(sheet.cell(row, 4), align="left")
        _body(sheet.cell(row, 5), align="left")
        sheet.cell(row, 5).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 30

    top = top + 3 + incidents
    _section(sheet, f"B{top}", "Renewal queue")
    rows = [("Renewals due from the as-of date", "=COUNTA(rn_subscription_id)"),
            ("High risk (failed card attempt or overdue)", '=COUNTIF(rn_risk_level,"high")'),
            ("Medium risk (due within 14 days)", '=COUNTIF(rn_risk_level,"medium")'),
            ("Failed card attempts on open renewals", "=SUM(rn_failed_attempts)")]
    for offset, (label, formula) in enumerate(rows, start=1):
        sheet.cell(top + offset, 2, label)
        sheet.cell(top + offset, 6, formula)
        sheet.merge_cells(start_row=top + offset, start_column=2, end_row=top + offset, end_column=5)
        _body(sheet.cell(top + offset, 2), align="left")
        _body(sheet.cell(top + offset, 6), COUNT, bold=True)
    sheet.sheet_properties.tabColor = BLUE


# --------------------------------------------------------------------------
# Audit, definitions, cover
# --------------------------------------------------------------------------

def _audit(book: Book, paid_campaigns: int) -> int:
    sheet = book.wb["Audit"]
    _header_band(sheet, "Audit", "Reconciliation checks over the data sheets. Every difference must be zero; "
                                 "they test source data, not constants.", 6)
    _table_header(sheet, 6, 2, ["Check", "Difference (cents or rows)", "Status"], [56, 22, 16])
    paid = ('SUMIFS(cp_{c},cp_medium,"paid_social")+SUMIFS(cp_{c},cp_medium,"paid_search")')
    checks = [
        ("Gross collected − refunds − net collected", "=Revenue!B2-Revenue!C2-Revenue!D2"),
        ("Bridge: CRM booked + adjustments − gross collected", "=SUM(Bridge!D2:D6)-Bridge!D7"),
        ("Bridge: gross collected + refunds − net collected", "=Bridge!D7+Bridge!D8-Bridge!D9"),
        ("Bridge net − revenue net", "=Bridge!D9-Revenue!D2"),
        ("Daily net cash − revenue net", "=SUM(dly_net_cash_cents)-Revenue!D2"),
        ("Attributed net cash (every payment, lead-creation) − revenue net", "=SUM(att_net_cash_cents)-Revenue!D2"),
        ("Attributed paid cash − platform warehouse cash",
         "=SUMPRODUCT(SUMIFS(att_net_cash_cents,att_campaign_id,dc_campaign_id)*dc_is_paid)"
         "-SUM(pl_warehouse_net_cash_cents)"),
        ("Daily spend − campaign spend", "=SUM(dly_spend_cents)-SUM(cp_spend_cents)"),
        ("Paid daily spend − campaign paid spend", "=SUM(pd_spend_cents)-(" + paid.format(c="spend_cents") + ")"),
        ("Paid daily spend − platform spend", "=SUM(pd_spend_cents)-SUM(pl_spend_cents)"),
        ("Paid daily leads − campaign paid leads", "=SUM(pd_leads)-(" + paid.format(c="leads") + ")"),
        ("Email delivered + bounces − sent", "=SUM(em_delivered)+SUM(em_bounces)-SUM(em_sends)"),
        ("Scorecard spend (all paid campaigns, all time) − paid daily spend",
         "=SUMPRODUCT(SUMIFS(pd_spend_cents,pd_campaign_id,'Campaign scorecard'!$B$7:$B$"
         f"{6 + paid_campaigns}))-SUM(pd_spend_cents)"),
        ("Paid daily rows whose campaign is missing from the campaign dimension",
         "=SUMPRODUCT(--ISNA(MATCH(pd_campaign_id,dc_campaign_id,0)))"),
    ]
    for offset, (label, formula) in enumerate(checks, start=7):
        sheet.cell(offset, 2, label)
        sheet.cell(offset, 3, formula)
        sheet.cell(offset, 4, f'=IF(C{offset}=0,"✓ Reconciled","✗ Investigate")')
        _body(sheet.cell(offset, 2), align="left")
        _body(sheet.cell(offset, 3), COUNT)
        _body(sheet.cell(offset, 4), align="left")
    last = 6 + len(checks)
    sheet.conditional_formatting.add(f"D7:D{last}", FormulaRule(formula=['LEFT(D7,1)="✗"'],
                                                                  font=Font(name=FONT, color=RED, bold=True)))
    sheet.conditional_formatting.add(f"D7:D{last}", FormulaRule(formula=['LEFT(D7,1)="✓"'],
                                                                  font=Font(name=FONT, color=GREEN, bold=True)))
    total = last + 2
    sheet.cell(total, 2, "All checks reconciled")
    sheet.cell(total, 3, f'=IF(SUMPRODUCT(ABS(C7:C{last}))=0,"Yes","No")')
    _body(sheet.cell(total, 2), bold=True, align="left")
    _body(sheet.cell(total, 3), bold=True)
    sheet.cell(total, 2).fill = sheet.cell(total, 3).fill = _fill(HEAD)
    book.name("AllChecksPass", f"Audit!$C${total}")
    sheet.sheet_properties.tabColor = GREEN
    return len(checks)


def _definitions(book: Book) -> None:
    sheet = book.wb["Definitions"]
    _header_band(sheet, "Metric definitions", "Copied from docs/metric-catalog.md at build time.", 4)
    _table_header(sheet, 5, 2, ["Metric", "Definition", "Grain / caveat"], [30, 70, 70])
    row = 6
    for line in CATALOG.read_text(encoding="utf-8").splitlines():
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if line.startswith("| ") and len(cells) == 3 and cells[0] != "Metric" and not cells[0].startswith("---"):
            for offset, value in enumerate(cells):
                cell = sheet.cell(row, 2 + offset, value.replace("`", ""))
                cell.font = _font(9, bold=offset == 0)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(bottom=_rule())
            row += 1
    sheet.sheet_properties.tabColor = "8A929C"


def _cover(book: Book, checks: int) -> None:
    sheet = book.wb["Cover"]
    _header_band(sheet, "ScaleLab GrowthOps · marketing and revenue workbook",
                 "Synthetic portfolio case: 15 months of generated data with planted incidents. No real company's "
                 "data.", 10)
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 90
    facts = [
        ("Data through", '=TEXT(AsOf,"d mmmm yyyy")'),
        ("Reporting window", '=Dashboard!E5&": "&TEXT(WinStart,"d mmm yyyy")&" – "&TEXT(WinEnd,"d mmm yyyy")'),
        ("Audit", f'=IF(AllChecksPass="Yes","All {checks} reconciliation checks are zero",'
                  '"A reconciliation check failed: see Audit")'),
    ]
    for offset, (label, formula) in enumerate(facts, start=5):
        sheet.cell(offset, 2, label).font = _font(10, color=MUTED)
        sheet.cell(offset, 3, formula).font = _font(11, bold=True, color=NAVY)
    _section(sheet, "B9", "Sheets")
    guide = [
        ("Dashboard", "Six KPIs for the window against the window before it, a written summary, monthly trend "
                      "and ROAS by platform. Change the window here."),
        ("Campaign scorecard", "Every paid campaign for the window: spend, CTR, CPL, MQL rate, cost per call, "
                               "cash and ROAS, with flags."),
        ("Email health", "Deliverability and engagement by sending domain and email type, with limit checks."),
        ("Revenue truth", "CRM bookings to net cash, step by step, and platform claims against warehouse cash."),
        ("Funnel and test", "Lead to renewal, and the CTA test decided on cash per visitor."),
        ("Data quality", "Tracking and migration checks against target, the short-link registry, incidents, "
                         "renewals."),
        ("Audit", "Reconciliation checks that must all be zero."),
        ("Definitions", "The governed metric catalog."),
        ("Calc", "How the window dates and chart series are derived."),
    ]
    for offset, (name, text) in enumerate(guide, start=10):
        cell = sheet.cell(offset, 2, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name=FONT, size=10, bold=True, color=BLUE, underline="single")
        sheet.cell(offset, 3, text).font = _font(10)
        sheet.cell(offset, 3).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[offset].height = 28
    after = 10 + len(guide) + 1
    _section(sheet, f"B{after}", "How to use it")
    notes = [
        "Yellow cells are the only inputs: the reporting window on the Dashboard, and the custom dates when the "
        "window is Custom. Every other cell is locked (no password) against accidental edits.",
        "Every report figure is a formula over the grey data sheets, which are the governed CSVs the Power BI "
        "model embeds. Formulas use named ranges: pd_spend_cents is the spend column of 'Paid daily'.",
        "The data sheets are Excel tables: filter and sort them freely. Rebuild the workbook with "
        "python -m growthops.export_excel.",
    ]
    for offset, text in enumerate(notes, start=after + 1):
        sheet.merge_cells(start_row=offset, start_column=2, end_row=offset, end_column=3)
        sheet.cell(offset, 2, "• " + text).font = _font(10)
        sheet.cell(offset, 2).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[offset].height = 30
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.tabColor = NAVY


# --------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------

def _months(first: date, last: date) -> list[date]:
    out, month = [], first.replace(day=1)
    while month <= last:
        out.append(month)
        month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
    return out


def build(output: Path = OUTPUT) -> Path:
    book = Book()
    book.wb.active.title = REPORT_SHEETS[0]
    for title in REPORT_SHEETS[1:]:
        book.wb.create_sheet(title)
    layout = {title: _data_sheet(book, title, source, prefix) for title, source, prefix in DATA_SHEETS}

    _, daily = _load("mart_growth_daily")
    _, email = _load("mart_email_performance")
    _, bridge = _load("mart_revenue_bridge")
    dim_header, dims = _load("dim_campaign")
    paid_campaigns = [(row[0], row[2]) for row in dims if row[dim_header.index("is_paid")]]
    book.paid_campaigns = len(paid_campaigns)
    domains = sorted({row[4] for row in email})
    types = sorted({row[2] for row in email})

    calc = _calc_sheet(book, _months(daily[0][0], daily[-1][0]), domains, bridge)
    _scorecard(book, paid_campaigns)
    _email(book, calc, types)
    _revenue(book, calc, bridge, layout["Platforms"]["rows"])
    _dashboard(book, calc)
    _funnel(book, layout["Funnel"]["rows"])
    _quality(book, layout["Quality"]["rows"], layout["Links"]["rows"], layout["Incidents"]["rows"])
    checks = _audit(book, len(paid_campaigns))
    _definitions(book)
    _cover(book, checks)

    for sheet in book.wb.worksheets:
        # Locked against accidental edits, no password; filters and sorting stay usable.
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.sort = False
        sheet.protection.formatColumns = False
    book.wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    book.wb.properties.creator = "GrowthOps OS"
    book.wb.properties.title = "ScaleLab GrowthOps workbook (synthetic)"
    book.wb.properties.created = book.wb.properties.modified = datetime(2026, 1, 1)
    book.wb.calculation.fullCalcOnLoad = True
    book.wb.save(output)
    _normalize_zip(output)
    return output


def _normalize_zip(path: Path) -> None:
    """Fix zip timestamps so an unchanged workbook rebuilds byte-for-byte (CI drift check)."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if name == "docProps/core.xml":  # openpyxl stamps save time into "modified"
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def recalculate(path: Path) -> None:
    """Open the workbook in Excel, recalculate and save, so it carries cached values.

    Windows with Excel only. A separate Excel instance is used so a workbook the user
    has open elsewhere is never touched.
    """
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        book = excel.Workbooks.Open(str(path.resolve()))
        excel.CalculateFull()
        book.Worksheets("Cover").Activate()
        book.Save()
        book.Close(False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--output", default=str(OUTPUT))
    parser.add_argument("--recalculate", action="store_true",
                        help="open the built workbook in Excel and save it with calculated values (Windows)")
    args = parser.parse_args()
    path = build(Path(args.output))
    if args.recalculate:
        recalculate(path)
    print(path)


if __name__ == "__main__":
    main()
